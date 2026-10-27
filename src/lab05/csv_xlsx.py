import csv
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    try:
        with open(csv_path, 'r', encoding='utf-8', newline='') as f:
            reader=csv.reader(f)
            rows=list(reader)
    except FileNotFoundError:
        raise FileNotFoundError("Файл не найден")
    if not rows:
        raise ValueError('CSV пустой')
    if all(len(row)==0 for row in rows):
        raise ValueError('CSV пустой')
    wb=Workbook()
    ws=wb.active
    ws.title='Sheet1'
    for row in rows:
        ws.append(row)
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len=0
        for cell in col:
            if cell.value is not None:
                max_len=max(max_len,len(str(cell.value)))
        ad_width=max(max_len, 8)
        col_letter=get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width=ad_width
    wb.save(xlsx_path)